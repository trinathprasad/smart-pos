import csv
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO

from flask import Blueprint, Response, flash, redirect, render_template, request, send_file, url_for
from sqlalchemy import and_, case, func, or_
from sqlalchemy.exc import IntegrityError

from ..extensions import db
from ..models import Category, Product, SaleItem, ShopSetting, StockMovement
from ..utils import LOCAL_TIMEZONE, format_indian_number, local_today, to_decimal, utc_bounds_for_local_date


products_bp = Blueprint("products", __name__, url_prefix="/products")


def _product_statistics():
    statistics = _inventory_statistics()
    return {
        "total_products": statistics["total_products"],
        "low_stock": statistics["low_stock"],
        "out_of_stock": statistics["out_of_stock"],
        "inventory_value": statistics["inventory_value"],
        "inventory_value_display": statistics["inventory_value_display"],
    }


def _inventory_statistics():
    total_products, low_stock, out_of_stock, inventory_value = (
        db.session.query(
            func.count(Product.id),
            func.coalesce(
                func.sum(case((and_(Product.stock_qty > 0, Product.stock_qty <= Product.low_stock_threshold), 1), else_=0)),
                0,
            ),
            func.coalesce(func.sum(case((Product.stock_qty <= 0, 1), else_=0)), 0),
            func.coalesce(func.sum(Product.purchase_price * Product.stock_qty), 0),
        )
        .filter(Product.is_active.is_(True))
        .one()
    )
    unit_totals = (
        db.session.query(Product.unit, func.sum(Product.stock_qty))
        .filter(Product.is_active.is_(True))
        .group_by(Product.unit)
        .order_by(Product.unit.asc())
        .all()
    )
    return {
        "total_products": total_products,
        "unit_totals": unit_totals,
        "low_stock": low_stock,
        "out_of_stock": out_of_stock,
        "inventory_value": inventory_value,
        "inventory_value_display": format_indian_number(inventory_value),
    }


def _inventory_filter_values():
    query = request.args.get("q", "").strip()
    category_filter = request.args.get("category", "").strip()
    stock_filter = request.args.get("filter", "").strip().lower()
    if stock_filter not in {"healthy", "low_stock", "out_of_stock"}:
        stock_filter = ""
    return query, category_filter, stock_filter


def _inventory_product_query(query="", category_filter="", stock_filter=""):
    product_query = (
        Product.query.outerjoin(Category, Product.category_id == Category.id)
        .filter(Product.is_active.is_(True))
    )
    if query:
        like = f"%{query}%"
        product_query = product_query.filter(or_(Product.name.ilike(like), Product.sku.ilike(like)))
    if category_filter.isdigit():
        product_query = product_query.filter(Product.category_id == int(category_filter))
    elif category_filter:
        product_query = product_query.filter(Product.category == category_filter)
    if stock_filter == "healthy":
        product_query = product_query.filter(Product.stock_qty > 0, Product.stock_qty > Product.low_stock_threshold)
    elif stock_filter == "low_stock":
        product_query = product_query.filter(Product.stock_qty > 0, Product.stock_qty <= Product.low_stock_threshold)
    elif stock_filter == "out_of_stock":
        product_query = product_query.filter(Product.stock_qty <= 0)
    return product_query


def _stock_status(product):
    if product.is_out_of_stock():
        return "Out of Stock"
    if product.is_low_stock():
        return "Low Stock"
    return "Healthy"


def _inventory_report_stats(products):
    inventory_value = sum(
        (product.purchase_price * product.stock_qty for product in products),
        start=Decimal("0.00"),
    )
    return {
        "total_products": len(products),
        "low_stock": sum(1 for product in products if product.is_low_stock()),
        "out_of_stock": sum(1 for product in products if product.is_out_of_stock()),
        "inventory_value": inventory_value,
        "inventory_value_display": format_indian_number(inventory_value),
    }


def _safe_spreadsheet_text(value):
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _inventory_export_rows(products):
    return [
        [
            _safe_spreadsheet_text(product.sku),
            _safe_spreadsheet_text(product.name),
            _safe_spreadsheet_text((product.category_record.name if product.category_record else product.category) or ""),
            _safe_spreadsheet_text(product.unit),
            product.purchase_price,
            product.selling_price,
            product.stock_qty,
            product.low_stock_threshold,
            _safe_spreadsheet_text(_stock_status(product)),
            product.purchase_price * product.stock_qty,
        ]
        for product in products
    ]


def _filtered_inventory_products():
    query, category_filter, stock_filter = _inventory_filter_values()
    products = _inventory_product_query(query, category_filter, stock_filter).order_by(Product.name.asc(), Product.sku.asc()).all()
    return products, query, category_filter, stock_filter


def _clean_barcode(value: str | None) -> str | None:
    barcode = (value or "").strip()
    return barcode or None


def _barcode_exists(barcode: str | None, product_id: int | None = None) -> bool:
    if not barcode:
        return False
    query = Product.query.filter(Product.barcode == barcode)
    if product_id is not None:
        query = query.filter(Product.id != product_id)
    return db.session.query(query.exists()).scalar()


def _active_categories():
    return Category.query.order_by(Category.name.asc()).all()


def _safe_return_path():
    return_to = request.values.get("return_to", "").strip()
    if return_to.startswith("/") and not return_to.startswith("//"):
        return return_to
    return None


def _parse_adjustment_quantity(raw_value):
    raw_value = (raw_value or "").strip()
    if not raw_value:
        return None, "Quantity is required."
    try:
        quantity = Decimal(raw_value)
    except (InvalidOperation, TypeError, ValueError):
        return None, "Quantity must be a valid decimal."
    if not quantity.is_finite() or quantity < 0:
        return None, "Quantity must be a non-negative decimal."
    if quantity.as_tuple().exponent < -2:
        return None, "Quantity can have at most 2 decimal places."
    return quantity, None


def _adjustment_form_context(product, return_to=None, form_data=None):
    return {
        "product": product,
        "return_to": return_to,
        "form_data": form_data or {},
    }


def _render_product_form(product=None, selected_category_id=None):
    return render_template(
        "products/form.html",
        product=product,
        categories=_active_categories(),
        selected_category_id=selected_category_id,
        return_to=_safe_return_path(),
    )


def _category_from_form():
    raw_category_id = request.form.get("category_id", "").strip()
    if not raw_category_id:
        return None, None
    try:
        category_id = int(raw_category_id)
    except ValueError:
        return None, "Choose a valid category."
    category = db.session.get(Category, category_id)
    if category is None:
        return None, "Choose a valid category."
    return category, None


def _required_price(field_name: str, required_message: str):
    raw_value = request.form.get(field_name)
    if raw_value is None or raw_value.strip() == "":
        return None, required_message
    return to_decimal(raw_value), None


@products_bp.route("/")
def index():
    query = request.args.get("q", "").strip()
    sku_filter = request.args.get("sku", "").strip()
    name_filter = request.args.get("name", "").strip()
    category_filter = request.args.get("category", "").strip()
    selected_category = None
    if category_filter:
        if category_filter.isdigit():
            selected_category = db.session.get(Category, int(category_filter))
        if selected_category is None:
            normalized_filter = category_filter.casefold()
            selected_category = Category.query.filter_by(normalized_name=normalized_filter).first()
        if selected_category:
            category_filter = selected_category.name
    stock_filter = request.args.get("filter", "").strip().lower()
    sort = request.args.get("sort", "name").strip().lower()
    direction = request.args.get("direction", "asc").strip().lower()
    if stock_filter not in {"healthy", "low_stock", "out_of_stock"}:
        stock_filter = ""

    sort_columns = {
        "sku": Product.sku,
        "name": Product.name,
        "category": Category.name,
        "stock": Product.stock_qty,
        "purchase_price": Product.purchase_price,
        "selling_price": Product.selling_price,
    }
    if sort not in sort_columns:
        sort = "name"
    if direction not in {"asc", "desc"}:
        direction = "asc"

    product_query = Product.query.outerjoin(Category, Product.category_id == Category.id).filter(Product.is_active.is_(True))
    if query:
        like = f"%{query}%"
        search_columns = [Product.name, Product.sku, Product.category, Category.name]
        barcode_column = getattr(Product, "barcode", None)
        if barcode_column is not None:
            search_columns.append(barcode_column)
        product_query = product_query.filter(or_(*(column.ilike(like) for column in search_columns)))
    if sku_filter:
        product_query = product_query.filter(Product.sku.ilike(f"%{sku_filter}%"))
    if name_filter:
        product_query = product_query.filter(Product.name.ilike(f"%{name_filter}%"))
    if selected_category:
        product_query = product_query.filter(Product.category_id == selected_category.id)
    elif category_filter:
        product_query = product_query.filter(Product.category == category_filter)
    if stock_filter == "healthy":
        product_query = product_query.filter(Product.stock_qty > 0, Product.stock_qty > Product.low_stock_threshold)
    elif stock_filter == "low_stock":
        product_query = product_query.filter(Product.stock_qty > 0, Product.stock_qty <= Product.low_stock_threshold)
    elif stock_filter == "out_of_stock":
        product_query = product_query.filter(Product.stock_qty <= 0)

    sort_column = sort_columns[sort]
    order_expression = sort_column.desc() if direction == "desc" else sort_column.asc()
    product_query = product_query.order_by(order_expression, Product.name.asc())

    products = product_query.all()
    categories = Category.query.join(Product, Product.category_id == Category.id).filter(Product.is_active.is_(True)).distinct().order_by(Category.name.asc()).all()
    return render_template(
        "products/index.html",
        products=products,
        product_stats=_product_statistics(),
        categories=categories,
        query=query,
        sku_filter=sku_filter,
        name_filter=name_filter,
        category_filter=category_filter,
        stock_filter=stock_filter,
        sort=sort,
        direction=direction,
    )


@products_bp.route("/inventory")
def inventory():
    products, query, category_filter, stock_filter = _filtered_inventory_products()
    categories = (
        Category.query.join(Product, Product.category_id == Category.id)
        .filter(Product.is_active.is_(True))
        .distinct()
        .order_by(Category.name.asc())
        .all()
    )
    matching_low_stock_count = _inventory_product_query(query, category_filter, "low_stock").count()
    matching_out_of_stock_count = _inventory_product_query(query, category_filter, "out_of_stock").count()
    return render_template(
        "products/inventory.html",
        products=products,
        inventory_stats=_inventory_statistics(),
        categories=categories,
        query=query,
        category_filter=category_filter,
        stock_filter=stock_filter,
        matching_product_count=len(products),
        matching_low_stock_count=matching_low_stock_count,
        matching_out_of_stock_count=matching_out_of_stock_count,
    )


def _inventory_timestamp():
    return datetime.now(LOCAL_TIMEZONE).strftime("%Y%m%d-%H%M%S")


@products_bp.route("/inventory/export.csv")
def inventory_export_csv():
    products, _, _, _ = _filtered_inventory_products()
    csv_buffer = StringIO(newline="")
    writer = csv.writer(csv_buffer)
    writer.writerow(
        [
            "SKU",
            "Product Name",
            "Category",
            "Unit",
            "Purchase Price",
            "Selling Price",
            "Current Stock",
            "Low-stock Alert",
            "Stock Status",
            "Inventory Purchase Value",
        ]
    )
    writer.writerows(_inventory_export_rows(products))
    filename = f"inventory-{_inventory_timestamp()}.csv"
    return Response(
        csv_buffer.getvalue().encode("utf-8-sig"),
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@products_bp.route("/inventory/export.xlsx")
def inventory_export_excel():
    products, _, _, _ = _filtered_inventory_products()
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return Response("Excel export is unavailable because openpyxl is not installed.", status=503, mimetype="text/plain")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory"
    headers = [
        "SKU",
        "Product Name",
        "Category",
        "Unit",
        "Purchase Price",
        "Selling Price",
        "Current Stock",
        "Low-stock Alert",
        "Stock Status",
        "Inventory Purchase Value",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A8EC2")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in _inventory_export_rows(products):
        sheet.append([float(value) if isinstance(value, Decimal) else value for value in row])
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = '₹#,##0.00'
    for row in sheet.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = '#,##0.00'
    for cell in sheet.iter_cols(min_col=10, max_col=10, min_row=2):
        for value_cell in cell:
            value_cell.number_format = '₹#,##0.00'
    widths = [18, 28, 20, 12, 16, 16, 16, 16, 18, 26]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"inventory-{_inventory_timestamp()}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@products_bp.route("/inventory/print")
def inventory_print():
    products, query, category_filter, stock_filter = _filtered_inventory_products()
    shop = ShopSetting.query.first()
    return render_template(
        "products/inventory_print.html",
        products=products,
        report_stats=_inventory_report_stats(products),
        shop=shop,
        generated_at=datetime.now(timezone.utc),
        query=query,
        category_filter=category_filter,
        stock_filter=stock_filter,
    )


@products_bp.route("/<int:product_id>/adjust-stock", methods=["GET", "POST"])
def adjust_stock(product_id):
    product = Product.query.filter(Product.id == product_id, Product.is_active.is_(True)).first()
    return_to = _safe_return_path()
    if product is None:
        flash("Archived or unavailable products cannot be adjusted.", "danger")
        return redirect(return_to or url_for("products.inventory"))

    if request.method == "POST":
        adjustment_type = request.form.get("adjustment_type", "").strip().lower()
        reason = request.form.get("reason", "").strip()
        quantity, quantity_error = _parse_adjustment_quantity(request.form.get("quantity"))
        form_data = request.form.to_dict(flat=True)
        errors = []
        if adjustment_type not in {"increase", "decrease", "set"}:
            errors.append("Choose a valid adjustment type.")
        if quantity_error:
            errors.append(quantity_error)
        if not reason:
            errors.append("An adjustment reason is required.")
        if adjustment_type in {"increase", "decrease"} and quantity is not None and quantity <= 0:
            errors.append("Increase and decrease quantities must be greater than zero.")

        if errors:
            for error in errors:
                flash(error, "danger")
            return render_template("products/stock_adjust.html", **_adjustment_form_context(product, return_to, form_data))

        previous_qty = Decimal(product.stock_qty)
        if adjustment_type == "increase":
            new_qty = previous_qty + quantity
        elif adjustment_type == "decrease":
            new_qty = previous_qty - quantity
        else:
            new_qty = quantity
        quantity_changed = new_qty - previous_qty
        if new_qty < 0:
            flash("The adjustment cannot reduce stock below zero.", "danger")
            return render_template("products/stock_adjust.html", **_adjustment_form_context(product, return_to, form_data))

        try:
            product.stock_qty = new_qty
            db.session.add(
                StockMovement(
                    product_id=product.id,
                    adjustment_type=adjustment_type,
                    previous_qty=previous_qty,
                    quantity_changed=quantity_changed,
                    new_qty=new_qty,
                    reason=reason,
                )
            )
            db.session.commit()
        except Exception:
            db.session.rollback()
            flash("Could not save the stock adjustment. No changes were made.", "danger")
            return render_template("products/stock_adjust.html", **_adjustment_form_context(product, return_to, form_data))

        flash(
            f"Stock adjusted successfully. New stock: {new_qty.normalize():f} {product.unit}.",
            "success",
        )
        return redirect(return_to or url_for("products.inventory"))

    return render_template("products/stock_adjust.html", **_adjustment_form_context(product, return_to))


@products_bp.route("/<int:product_id>/stock-history")
def stock_history(product_id):
    product = Product.query.get_or_404(product_id)
    date_range = request.args.get("date_range", "").strip().lower()
    adjustment_type = request.args.get("adjustment_type", "").strip().lower()
    reason_query = request.args.get("reason", "").strip()
    start_date_value = request.args.get("start_date", "").strip()
    end_date_value = request.args.get("end_date", "").strip()
    allowed_date_ranges = {"today", "last_7_days", "last_30_days", "custom"}
    allowed_adjustment_types = {"increase", "decrease", "set"}
    if date_range not in allowed_date_ranges:
        date_range = ""
    if adjustment_type not in allowed_adjustment_types:
        adjustment_type = ""

    today = local_today()
    start_date = None
    end_date = None
    filter_errors = []
    if date_range == "today":
        start_date = end_date = today
    elif date_range == "last_7_days":
        start_date, end_date = today - timedelta(days=6), today
    elif date_range == "last_30_days":
        start_date, end_date = today - timedelta(days=29), today
    elif date_range == "custom":
        try:
            start_date = date.fromisoformat(start_date_value) if start_date_value else None
            end_date = date.fromisoformat(end_date_value) if end_date_value else None
        except ValueError:
            filter_errors.append("Enter valid custom start and end dates.")
        if start_date is None or end_date is None:
            filter_errors.append("Custom date range requires both a start date and an end date.")

    if start_date and end_date:
        if start_date > end_date:
            filter_errors.append("Start date cannot be later than end date.")
        if start_date > today or end_date > today:
            filter_errors.append("Future dates cannot be included in stock history.")

    movements = []
    if not filter_errors:
        movement_query = StockMovement.query.filter(StockMovement.product_id == product.id)
        if start_date and end_date:
            start_datetime, _ = utc_bounds_for_local_date(start_date)
            _, end_datetime = utc_bounds_for_local_date(end_date)
            movement_query = movement_query.filter(
                StockMovement.created_at >= start_datetime,
                StockMovement.created_at <= end_datetime,
            )
        else:
            _, today_end_datetime = utc_bounds_for_local_date(today)
            movement_query = movement_query.filter(StockMovement.created_at <= today_end_datetime)
        if adjustment_type:
            movement_query = movement_query.filter(StockMovement.adjustment_type == adjustment_type)
        if reason_query:
            movement_query = movement_query.filter(StockMovement.reason.ilike(f"%{reason_query}%"))
        movements = movement_query.order_by(StockMovement.created_at.desc(), StockMovement.id.desc()).all()
    for error in filter_errors:
        flash(error, "danger")

    return render_template(
        "products/stock_history.html",
        product=product,
        movements=movements,
        return_to=_safe_return_path(),
        date_range=date_range,
        adjustment_type=adjustment_type,
        reason_query=reason_query,
        start_date_value=start_date_value,
        end_date_value=end_date_value,
        today_date=today.isoformat(),
    )


@products_bp.route("/export")
def export():
    products = Product.query.filter_by(is_active=True).order_by(Product.name.asc()).all()
    csv_buffer = StringIO()
    writer = csv.writer(csv_buffer)
    writer.writerow(
        ["SKU", "Name", "Category", "Unit", "Purchase Price", "Selling Price", "Stock Qty", "Low Stock Threshold"]
    )
    for product in products:
        writer.writerow(
            [
                product.sku,
                product.name,
                (product.category_record.name if product.category_record else product.category) or "",
                product.unit,
                product.purchase_price,
                product.selling_price,
                product.stock_qty,
                product.low_stock_threshold,
            ]
        )
    return Response(
        csv_buffer.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=products.csv"},
    )


@products_bp.route("/new", methods=["GET", "POST"])
def create():
    if request.method == "POST":
        category, category_error = _category_from_form()
        if category_error:
            flash(category_error, "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))
        sku = request.form.get("sku", "").strip()
        name = request.form.get("name", "").strip()
        barcode = _clean_barcode(request.form.get("barcode"))
        if not sku or not name:
            flash("SKU and product name are required.", "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))
        if _barcode_exists(barcode):
            flash("Barcode already exists for another product.", "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))
        purchase_price, price_error = _required_price("purchase_price", "Purchase price is required.")
        if price_error:
            flash(price_error, "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))
        selling_price, price_error = _required_price("selling_price", "Selling price is required.")
        if price_error:
            flash(price_error, "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))
        stock_qty = to_decimal(request.form.get("stock_qty"))
        low_stock_threshold = to_decimal(request.form.get("low_stock_threshold"))
        if purchase_price < 0:
            flash("Purchase price cannot be negative.", "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))
        if selling_price < 0:
            flash("Selling price cannot be negative.", "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))
        if stock_qty < 0 or low_stock_threshold < 0:
            flash("Stock and low stock values cannot be negative.", "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))

        product = Product(
            sku=sku,
            barcode=barcode,
            name=name,
            category=category.name if category else None,
            category_record=category,
            brand=request.form.get("brand", "").strip() or None,
            unit=request.form.get("unit", "pcs").strip() or "pcs",
            description=request.form.get("description", "").strip() or None,
            purchase_price=purchase_price,
            selling_price=selling_price,
            stock_qty=stock_qty,
            low_stock_threshold=low_stock_threshold,
        )
        try:
            db.session.add(product)
            db.session.commit()
        except IntegrityError as exc:
            db.session.rollback()
            error_text = str(exc.orig).lower()
            if "barcode" in error_text:
                flash("Barcode already exists for another product.", "danger")
            elif "sku" in error_text or "unique constraint failed" in error_text:
                flash("SKU must be unique. This code already exists.", "danger")
            else:
                flash(f"Could not save product: {exc.orig}", "danger")
            return _render_product_form(None, request.form.get("category_id", type=int))
        flash("Product added successfully.", "success")
        return redirect(_safe_return_path() or url_for("products.index"))
    selected_category_id = request.args.get("category_id", type=int)
    return _render_product_form(None, selected_category_id)


@products_bp.route("/<int:product_id>/edit", methods=["GET", "POST"])
def edit(product_id):
    product = Product.query.get_or_404(product_id)
    if request.method == "POST":
        category, category_error = _category_from_form()
        if category_error:
            flash(category_error, "danger")
            return _render_product_form(product, request.form.get("category_id", type=int))
        barcode = _clean_barcode(request.form.get("barcode"))
        if _barcode_exists(barcode, product.id):
            flash("Barcode already exists for another product.", "danger")
            return _render_product_form(product, request.form.get("category_id", type=int))
        purchase_price, price_error = _required_price("purchase_price", "Purchase price is required.")
        if price_error:
            flash(price_error, "danger")
            return _render_product_form(product, request.form.get("category_id", type=int))
        selling_price, price_error = _required_price("selling_price", "Selling price is required.")
        if price_error:
            flash(price_error, "danger")
            return _render_product_form(product, request.form.get("category_id", type=int))
        stock_qty = to_decimal(request.form.get("stock_qty"))
        low_stock_threshold = to_decimal(request.form.get("low_stock_threshold"))
        if purchase_price < 0:
            flash("Purchase price cannot be negative.", "danger")
            return _render_product_form(product, request.form.get("category_id", type=int))
        if selling_price < 0:
            flash("Selling price cannot be negative.", "danger")
            return _render_product_form(product, request.form.get("category_id", type=int))
        if stock_qty < 0 or low_stock_threshold < 0:
            flash("Stock and low stock values cannot be negative.", "danger")
            return _render_product_form(product, request.form.get("category_id", type=int))
        product.sku = request.form.get("sku", "").strip()
        product.barcode = barcode
        product.name = request.form.get("name", "").strip()
        product.category_record = category
        product.category = category.name if category else None
        product.brand = request.form.get("brand", "").strip() or None
        product.unit = request.form.get("unit", "pcs").strip() or "pcs"
        product.description = request.form.get("description", "").strip() or None
        product.purchase_price = purchase_price
        product.selling_price = selling_price
        product.stock_qty = stock_qty
        product.low_stock_threshold = low_stock_threshold
        try:
            db.session.commit()
        except IntegrityError as exc:
            db.session.rollback()
            error_text = str(exc.orig).lower()
            if "barcode" in error_text:
                flash("Barcode already exists for another product.", "danger")
            elif "sku" in error_text or "unique constraint failed" in error_text:
                flash("SKU must be unique. This code already exists.", "danger")
            else:
                flash(f"Could not update product: {exc.orig}", "danger")
            return _render_product_form(product, request.form.get("category_id", type=int))
        flash("Product updated successfully.", "success")
        return redirect(url_for("products.index"))
    return _render_product_form(product, request.form.get("category_id", type=int))


@products_bp.route("/<int:product_id>/delete", methods=["POST"])
def delete(product_id):
    product = Product.query.get_or_404(product_id)
    sale_item_count = SaleItem.query.filter_by(product_id=product.id).count()
    if sale_item_count:
        product.is_active = False
        product.sku = f"{product.sku}-ARCHIVED-{product.id}"
        db.session.commit()
        flash(
            f"{product.name} is used in {sale_item_count} sale item(s), so it was archived instead.",
            "info",
        )
        return redirect(url_for("products.index"))

    db.session.delete(product)
    db.session.commit()
    flash("Product deleted.", "info")
    return redirect(url_for("products.index"))
